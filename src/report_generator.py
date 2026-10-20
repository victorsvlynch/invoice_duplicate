import io
from typing import Dict, Any
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_excel_report(
    summary_df: pd.DataFrame,
    details_df: pd.DataFrame,
    stats: Dict[str, Any],
) -> bytes:
    """
    Generates a beautifully styled Excel workbook report in memory and returns bytes.
    Sheets:
    1. Resumen: Overview stats + summary table
    2. Detalle_Duplicados: Complete list of matched rows with Excel data
    3. No_Encontrados: Invoices that were not found in Excel
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    title_font = Font(name="Calibri", size=14, bold=True, color="1F497D")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="595959")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=10, bold=True)
    normal_font = Font(name="Calibri", size=10)

    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    sub_header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    stat_label_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    stat_val_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    dup_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")  # soft orange/red
    not_found_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")  # soft green
    multi_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # soft yellow

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    # -------------------------------------------------------------------------
    # SHEET 1: RESUMEN
    # -------------------------------------------------------------------------
    ws_resumen = wb.create_sheet(title="Resumen")
    ws_resumen.views.sheetView[0].showGridLines = True

    # Title
    ws_resumen.cell(row=1, column=1, value="Reporte de Verificación de Facturas Duplicadas").font = title_font
    ws_resumen.cell(row=2, column=1, value="Sistema de Comparación con Columna P de Excel").font = subtitle_font

    # Stats block
    stat_rows = [
        ("Total Items Consultados", stats.get("total_queries", 0)),
        ("Items Únicos Consultados", stats.get("unique_queries", 0)),
        ("Encontrados / Duplicados en Excel", stats.get("found_count", 0)),
        ("No Encontrados en Excel", stats.get("not_found_count", 0)),
        ("Con Múltiples Coincidencias", stats.get("multiple_matches_count", 0)),
        ("Repetidos en Lista Pegada", stats.get("internal_duplicates", 0)),
    ]

    for i, (label, val) in enumerate(stat_rows, start=4):
        c1 = ws_resumen.cell(row=i, column=1, value=label)
        c1.font = bold_font
        c1.fill = stat_label_fill
        c1.border = thin_border

        c2 = ws_resumen.cell(row=i, column=2, value=val)
        c2.font = bold_font
        c2.fill = stat_val_fill
        c2.alignment = center_align
        c2.border = thin_border

    # Table Header
    start_row = 11
    if not summary_df.empty:
        cols = list(summary_df.columns)
        for col_idx, col_name in enumerate(cols, start=1):
            cell = ws_resumen.cell(row=start_row, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        # Table Rows
        for r_idx, row_data in enumerate(summary_df.itertuples(index=False), start=start_row + 1):
            status = row_data[summary_df.columns.get_loc("Estado en Excel")]
            for col_idx, val in enumerate(row_data, start=1):
                cell = ws_resumen.cell(row=r_idx, column=col_idx, value=val)
                cell.font = normal_font
                cell.border = thin_border

                if "DUPLICADO" in str(status):
                    if "veces" in str(status):
                        cell.fill = multi_fill
                    else:
                        cell.fill = dup_fill
                else:
                    cell.fill = not_found_fill

                if col_idx in [1, 3, 4, 6, 7]:
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align

        # Auto column widths
        for col in ws_resumen.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_resumen.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # -------------------------------------------------------------------------
    # SHEET 2: DETALLE DUPLICADOS
    # -------------------------------------------------------------------------
    ws_detalles = wb.create_sheet(title="Detalle_Duplicados")
    ws_detalles.views.sheetView[0].showGridLines = True

    ws_detalles.cell(row=1, column=1, value="Detalle de Coincidencias Encontradas en Excel").font = title_font

    if not details_df.empty:
        cols = list(details_df.columns)
        for col_idx, col_name in enumerate(cols, start=1):
            cell = ws_detalles.cell(row=3, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = sub_header_fill
            cell.alignment = center_align
            cell.border = thin_border

        for r_idx, row_data in enumerate(details_df.itertuples(index=False), start=4):
            for col_idx, val in enumerate(row_data, start=1):
                cell = ws_detalles.cell(row=r_idx, column=col_idx, value=val)
                cell.font = normal_font
                cell.border = thin_border
                if col_idx in [2, 4]:
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align

        for col in ws_detalles.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_detalles.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)
    else:
        ws_detalles.cell(row=3, column=1, value="No se encontraron duplicados.").font = normal_font

    # -------------------------------------------------------------------------
    # SHEET 3: NO ENCONTRADOS
    # -------------------------------------------------------------------------
    ws_not_found = wb.create_sheet(title="No_Encontrados")
    ws_not_found.views.sheetView[0].showGridLines = True
    ws_not_found.cell(row=1, column=1, value="Facturas / Items No Encontrados en Excel").font = title_font

    if not summary_df.empty:
        not_found_df = summary_df[summary_df["Estado en Excel"] == "NO ENCONTRADO"]
        if not not_found_df.empty:
            cols = ["N°", "Texto Consultado", "Serie Detectada", "Correlativo Ceros", "Número Sin Ceros"]
            for col_idx, col_name in enumerate(cols, start=1):
                cell = ws_not_found.cell(row=3, column=col_idx, value=col_name)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border

            for r_idx, (_, r_data) in enumerate(not_found_df[cols].iterrows(), start=4):
                for col_idx, val in enumerate(r_data, start=1):
                    cell = ws_not_found.cell(row=r_idx, column=col_idx, value=val)
                    cell.font = normal_font
                    cell.border = thin_border
                    cell.fill = not_found_fill
                    if col_idx in [1, 3, 4, 5]:
                        cell.alignment = center_align
                    else:
                        cell.alignment = left_align

            for col in ws_not_found.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws_not_found.column_dimensions[col_letter].width = max(max_len + 3, 12)
        else:
            ws_not_found.cell(row=3, column=1, value="Todos los items consultados fueron encontrados.").font = normal_font

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
